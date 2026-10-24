"""
Fetch Sheet Resources
Downloads the resources linked from the "Resources" column of the video
resources Google Sheet (all quarter tabs) so the prompts can be reviewed:
- PDF guides, with their text extracted
- prompt-style text files (.md, .txt, extensionless files and Google Docs)
Byte-identical duplicates are kept once, with every linking video recorded.

Requirements: pip install requests openpyxl, plus the pdftotext CLI
(poppler-utils on Linux, `brew install poppler` on Mac).
The sheet and its Drive folders must be shared as "Anyone with the link".

Usage: python prompts/fetch_sheet_resources.py [output_dir]
"""
import hashlib
import html
import io
import json
import re
import subprocess
import sys
import time
from pathlib import Path

import requests
from openpyxl import load_workbook

SHEET_ID = "1h5k7vAAXl5_f4VtguDkHwCCy-S-dd7RfG3-rUiMxfPw"
FOLDER_ENTRY = re.compile(
    r'<div class="flip-entry" id="entry-([^"]+)".*?<a href="([^"]+)".*?'
    r'<div class="flip-entry-title">([^<]*)</div>',
    re.S,
)


def get(url):
    """GET with retries (2s, 4s, 8s backoff)"""
    for attempt in range(4):
        try:
            resp = requests.get(url, timeout=120)
            resp.raise_for_status()
            return resp
        except requests.RequestException:
            if attempt == 3:
                raise
            time.sleep(2 ** (attempt + 1))


def read_resource_links():
    """Return one dict per video row with its Resources link (if any)"""
    xlsx = get(f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx").content
    rows = []
    for ws in load_workbook(io.BytesIO(xlsx)).worksheets:
        header = [c.value for c in ws[1]]
        if "Resources" not in header:
            continue
        col = {name: header.index(name) + 1 for name in ("Publish Date", "Title", "Resources")}
        for r in range(2, ws.max_row + 1):
            title = ws.cell(r, col["Title"]).value
            cell = ws.cell(r, col["Resources"])
            if not title:
                continue
            # Links are stored either as HYPERLINK() formulas or as cell hyperlinks
            match = re.search(r'HYPERLINK\("([^"]+)"', str(cell.value or ""))
            url = match.group(1) if match else (cell.hyperlink.target if cell.hyperlink else None)
            rows.append({
                "quarter": ws.title,
                "row": r,
                "publish_date": str(ws.cell(r, col["Publish Date"]).value or ""),
                "title": title,
                "resources_link": url,
            })
    return rows


def list_drive_folder(folder_id, path="", seen=None):
    """Recursively list files in a public Drive folder"""
    seen = seen if seen is not None else set()
    if folder_id in seen:
        return []
    seen.add(folder_id)
    page = get(f"https://drive.google.com/embeddedfolderview?id={folder_id}").text
    files = []
    for entry_id, href, name in FOLDER_ENTRY.findall(page):
        href, name = html.unescape(href), html.unescape(name)
        if "/folders/" in href:
            files += list_drive_folder(entry_id, f"{path}{name}/", seen)
        else:
            files.append({"id": entry_id, "name": name, "path": path + name,
                          "google_doc": "docs.google.com/document" in href})
    return files


def kind_of(file):
    """Classify a Drive file as 'pdf', 'text' or None (skipped)"""
    name = file["name"].lower()
    if name.endswith(".pdf"):
        return "pdf"
    if file["google_doc"] or name.endswith((".md", ".txt")) or "." not in name:
        return "text"
    return None


def download(file):
    if file["google_doc"]:
        return get(f"https://docs.google.com/document/d/{file['id']}/export?format=txt").content
    return get(f"https://drive.usercontent.google.com/download?id={file['id']}&export=download&confirm=t").content


def main():
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "outputs/sheet_resources")
    for sub in ("pdf", "text", "files"):
        (out / sub).mkdir(parents=True, exist_ok=True)

    rows = read_resource_links()
    print(f"Sheet rows: {len(rows)}, with resources link: {sum(1 for r in rows if r['resources_link'])}")

    # Collect files from Drive folders and direct Drive file links
    found = {}
    for row in rows:
        url = row["resources_link"] or ""
        if "drive/folders/" in url:
            files = list_drive_folder(re.search(r"folders/([\w-]+)", url).group(1))
        elif "drive.google.com/file/d/" in url:
            file_id = re.search(r"file/d/([\w-]+)", url).group(1)
            files = [{"id": file_id, "name": f"{file_id}.pdf", "path": f"{file_id}.pdf", "google_doc": False}]
        else:
            continue
        for f in files:
            if kind_of(f):
                found.setdefault(f["id"], {**f, "kind": kind_of(f), "videos": []})["videos"].append(row)

    # Download, drop byte-identical duplicates, extract PDF text
    manifest, by_hash = [], {}
    for f in found.values():
        data = download(f)
        if f["kind"] == "pdf" and not data.startswith(b"%PDF"):
            continue
        digest = hashlib.md5(data).hexdigest()
        if digest in by_hash:
            by_hash[digest]["videos"] += f["videos"]
            continue
        entry = {**f, "md5": digest, "drive_link": f"https://drive.google.com/file/d/{f['id']}/view"}
        if f["kind"] == "pdf":
            pdf_path = out / "pdf" / f"{f['id']}.pdf"
            pdf_path.write_bytes(data)
            entry["text_file"] = str(out / "text" / f"{f['id']}.txt")
            subprocess.run(["pdftotext", "-layout", str(pdf_path), entry["text_file"]], check=True)
        else:
            entry["text_file"] = str(out / "files" / f"{f['id']}.txt")
            Path(entry["text_file"]).write_bytes(data)
        by_hash[digest] = entry
        manifest.append(entry)

    (out / "manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False))
    pdfs = sum(1 for e in manifest if e["kind"] == "pdf")
    print(f"Unique PDFs: {pdfs}, unique text files: {len(manifest) - pdfs} -> {out / 'manifest.json'}")


if __name__ == "__main__":
    main()
