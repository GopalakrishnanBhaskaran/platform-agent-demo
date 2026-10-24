"""
Fetch Sheet PDFs
Downloads every PDF linked from the "Resources" column of the video resources
Google Sheet (all quarter tabs), removes duplicates, and extracts their text so
the prompts can be reviewed.

Requirements: pip install requests openpyxl, plus the pdftotext CLI
(poppler-utils on Linux, `brew install poppler` on Mac).
The sheet and its Drive folders must be shared as "Anyone with the link".

Usage: python prompts/fetch_sheet_pdfs.py [output_dir]
"""
import hashlib
import html
import io
import json
import re
import subprocess
import sys
import time
from pathlib import Path

import requests
from openpyxl import load_workbook

SHEET_ID = "1h5k7vAAXl5_f4VtguDkHwCCy-S-dd7RfG3-rUiMxfPw"
FOLDER_ENTRY = re.compile(
    r'<div class="flip-entry" id="entry-([^"]+)".*?<a href="([^"]+)".*?'
    r'<div class="flip-entry-title">([^<]*)</div>',
    re.S,
)


def get(url):
    """GET with retries (2s, 4s, 8s backoff)"""
    for attempt in range(4):
        try:
            resp = requests.get(url, timeout=120)
            resp.raise_for_status()
            return resp
        except requests.RequestException:
            if attempt == 3:
                raise
            time.sleep(2 ** (attempt + 1))


def read_resource_links():
    """Return one dict per video row with its Resources link (if any)"""
    xlsx = get(f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx").content
    rows = []
    for ws in load_workbook(io.BytesIO(xlsx)).worksheets:
        header = [c.value for c in ws[1]]
        if "Resources" not in header:
            continue
        col = {name: header.index(name) + 1 for name in ("Publish Date", "Title", "Resources")}
        for r in range(2, ws.max_row + 1):
            title = ws.cell(r, col["Title"]).value
            cell = ws.cell(r, col["Resources"])
            if not title:
                continue
            # Links are stored either as HYPERLINK() formulas or as cell hyperlinks
            match = re.search(r'HYPERLINK\("([^"]+)"', str(cell.value or ""))
            url = match.group(1) if match else (cell.hyperlink.target if cell.hyperlink else None)
            rows.append({
                "quarter": ws.title,
                "row": r,
                "publish_date": str(ws.cell(r, col["Publish Date"]).value or ""),
                "title": title,
                "resources_link": url,
            })
    return rows


def list_drive_folder(folder_id, seen=None):
    """Recursively list files in a public Drive folder"""
    seen = seen if seen is not None else set()
    if folder_id in seen:
        return []
    seen.add(folder_id)
    page = get(f"https://drive.google.com/embeddedfolderview?id={folder_id}").text
    files = []
    for entry_id, href, name in FOLDER_ENTRY.findall(page):
        if "/folders/" in html.unescape(href):
            files += list_drive_folder(entry_id, seen)
        else:
            files.append({"id": entry_id, "name": html.unescape(name)})
    return files


def main():
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "outputs/sheet_pdfs")
    (out / "pdf").mkdir(parents=True, exist_ok=True)
    (out / "text").mkdir(parents=True, exist_ok=True)

    rows = read_resource_links()
    print(f"Sheet rows: {len(rows)}, with resources link: {sum(1 for r in rows if r['resources_link'])}")

    # Collect PDF file IDs from Drive folders and direct Drive file links
    pdfs = {}
    for row in rows:
        url = row["resources_link"] or ""
        if "drive/folders/" in url:
            files = list_drive_folder(re.search(r"folders/([\w-]+)", url).group(1))
            candidates = [f for f in files if f["name"].lower().endswith(".pdf")]
        elif "drive.google.com/file/d/" in url:
            candidates = [{"id": re.search(r"file/d/([\w-]+)", url).group(1), "name": None}]
        else:
            continue
        for f in candidates:
            pdfs.setdefault(f["id"], {"id": f["id"], "name": f["name"], "videos": []})["videos"].append(row)

    # Download, keep PDFs only, and drop byte-identical duplicates
    manifest, by_hash = [], {}
    for pdf in pdfs.values():
        data = get(f"https://drive.usercontent.google.com/download?id={pdf['id']}&export=download&confirm=t").content
        if not data.startswith(b"%PDF"):
            continue
        digest = hashlib.md5(data).hexdigest()
        if digest in by_hash:
            by_hash[digest]["videos"] += pdf["videos"]
            continue
        path = out / "pdf" / f"{pdf['id']}.pdf"
        path.write_bytes(data)
        text_path = out / "text" / f"{pdf['id']}.txt"
        subprocess.run(["pdftotext", "-layout", str(path), str(text_path)], check=True)
        entry = {**pdf, "md5": digest, "pdf_link": f"https://drive.google.com/file/d/{pdf['id']}/view",
                 "text_file": str(text_path)}
        by_hash[digest] = entry
        manifest.append(entry)

    (out / "manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False))
    print(f"Unique PDFs: {len(manifest)} -> {out / 'manifest.json'}")


if __name__ == "__main__":
    main()
